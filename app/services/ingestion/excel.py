# app/services/ingestion/excel.py
from __future__ import annotations

import io
from typing import TYPE_CHECKING, Any

from openpyxl import load_workbook

from app.services.ingestion.base import (
    IngestionError,
    IngestionResult,
    ParsedRow,
    SourceMetadata,
)
from app.services.storage.local import storage

if TYPE_CHECKING:
    from app.db.models.dataset import DataSource


def _coerce(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


class ExcelConnector:
    source_type = "excel"

    async def _load(self, *, source: DataSource, sheet_name: str | None = None):
        storage_key = source.config.get("storage_key")
        if not storage_key:
            raise IngestionError("source has no storage_key in config")
        raw = await storage.get(key=storage_key)
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:
            raise IngestionError(f"could not open Excel file: {exc}") from exc

        sheet = sheet_name or source.config.get("sheet_name")
        if sheet is None:
            ws = wb[wb.sheetnames[0]]
        else:
            if sheet not in wb.sheetnames:
                wb.close()
                raise IngestionError(f"sheet not found: {sheet}")
            ws = wb[sheet]
        return wb, ws

    async def inspect(self, *, source: DataSource) -> SourceMetadata:
        wb, ws = await self._load(source=source)
        try:
            it = ws.iter_rows(values_only=True)
            header_row = next(it, None)
            if header_row is None:
                raise IngestionError("Excel sheet is empty")
            header = [str(c) if c is not None else "" for c in header_row]

            sample_rows: list[dict[str, Any]] = []
            for _ in range(5):
                row = next(it, None)
                if row is None:
                    break
                sample_rows.append({
                    h: _coerce(v) for h, v in zip(header, row, strict=False)
                })
            return SourceMetadata(
                columns=header,
                sample_rows=sample_rows,
                sheet_names=list(wb.sheetnames),
            )
        finally:
            wb.close()

    async def ingest(self, *, source: DataSource) -> IngestionResult:
        wb, ws = await self._load(source=source)
        try:
            it = ws.iter_rows(values_only=True)
            header_row = next(it, None)
            if header_row is None:
                raise IngestionError("Excel sheet is empty")
            header = [str(c) if c is not None else "" for c in header_row]

            rows: list[ParsedRow] = []
            errors: list[dict[str, Any]] = []

            for idx, row in enumerate(it, start=2):
                if row is None:
                    continue
                data = {h: _coerce(v) for h, v in zip(header, row, strict=False)}
                if all(v is None or v == "" for v in data.values()):
                    continue
                if len(row) != len(header):
                    errors.append({
                        "row": idx,
                        "error": "column count mismatch",
                        "expected": len(header),
                        "got": len(row),
                    })
                    continue
                rows.append(ParsedRow(row_number=idx, data=data))

            return IngestionResult(
                rows=rows,
                metadata=SourceMetadata(
                    columns=header,
                    sheet_names=list(wb.sheetnames),
                ),
                errors=errors,
            )
        finally:
            wb.close()